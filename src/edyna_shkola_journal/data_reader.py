from openpyxl import load_workbook
from src.edyna_shkola_journal.models import Lesson, Semester, Theme, LessonXlsxRow
import pathlib


from collections import defaultdict


def group_lessons(rows: list[LessonXlsxRow]) -> list[Semester]:
    grouped = defaultdict(lambda: defaultdict(list))

    for row in rows:
        grouped[row.semester_number][row.theme_title].append(
            Lesson(
                topic=row.lesson_topic,
                notes=row.notes or None,
                homework=row.homework or None,
            )
        )

    return [
        Semester(
            semester=semester_number,
            themes=[
                Theme(
                    title=theme_title,
                    lessons=lessons,
                )
                for theme_title, lessons in themes.items()
            ],
        )
        for semester_number, themes in grouped.items()
    ]


def get_semesters(path: pathlib.Path) -> list[Semester]:
    workbook = load_workbook(path.as_posix(), data_only=True)
    sheet = workbook["теми"]

    xlsx_rows: list[LessonXlsxRow] = []

    for row in sheet.iter_rows(min_row=6, values_only=True):
        semester_number = row[1]  # B
        theme_title = row[4]  # E
        lesson_topic = row[7]  # H
        notes = row[8]  # I
        homework = row[9]  # J

        if lesson_topic is None:
            continue

        xlsx_rows.append(
            LessonXlsxRow(
                semester_number=int(str(semester_number)),
                lesson_topic=str(lesson_topic),
                homework=str(homework),
                notes=str(notes),
                theme_title=str(theme_title),
            )
        )

    return group_lessons(xlsx_rows)
